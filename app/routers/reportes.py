import io
from datetime import datetime, timedelta, timezone
from typing import Optional
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, field_validator
from sqlalchemy import func
from sqlalchemy.orm import Session

from app import models
from app.database import get_db
from app.dependencies import get_current_user, get_current_user_admin
from app.services.venta_service import obtener_ventas_empresa

router = APIRouter(prefix="/reportes", tags=["Reportes y Simulador Comercial"])

_HEADER_FILL = PatternFill("solid", fgColor="1E40AF")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


class SimulacionComercialRequest(BaseModel):
    nombre_escenario: str = "Escenario Simulado"
    precio_venta_actual: Optional[float] = None
    costo_compra_actual: Optional[float] = None
    volumen_ventas_actual: Optional[float] = None
    gastos_fijos_actuales: float = 0.0
    gastos_variables_pct_actual: float = 0.0

    # Variables simuladas o directas desde el frontend interactivo
    precio_venta_simulado: Optional[float] = None
    costo_compra_simulado: Optional[float] = None
    volumen_ventas_simulado: Optional[float] = None
    descuento_promocion_pct: float = 0.0
    gastos_fijos_simulados: Optional[float] = None
    gastos_variables_pct_simulado: Optional[float] = None

    # Aliases directos (Simulador.jsx / tests)
    precio_venta: Optional[float] = None
    precio_compra: Optional[float] = None
    volumen_unidades: Optional[float] = None
    gastos_fijos: Optional[float] = None
    gasto_variable_unitario: float = 0.0
    descuento_porcentaje: Optional[float] = None
    precio_venta_base: Optional[float] = None
    costo_compra_base: Optional[float] = None
    volumen_base: Optional[float] = None


def _calcular_escenario(
    precio_lista: float,
    descuento_pct: float,
    costo_unitario: float,
    volumen_unidades: float,
    gastos_fijos: float,
    gastos_variables_pct: float = 0.0,
    gasto_variable_unitario_fijo: float = 0.0,
) -> dict:
    import math
    descuento_frac = max(0.0, min(100.0, descuento_pct)) / 100.0
    gasto_var_frac = max(0.0, min(100.0, gastos_variables_pct)) / 100.0

    precio_efectivo = round(precio_lista * (1.0 - descuento_frac), 2)
    gasto_variable_unitario = round(
        (precio_efectivo * gasto_var_frac) + float(gasto_variable_unitario_fijo or 0.0),
        2,
    )
    costo_total_unitario = round(costo_unitario + gasto_variable_unitario, 2)

    margen_unitario = round(precio_efectivo - costo_total_unitario, 2)
    margen_porcentual = (
        round((margen_unitario / precio_efectivo) * 100.0, 2)
        if precio_efectivo > 0
        else 0.0
    )

    ingresos_totales = round(precio_efectivo * volumen_unidades, 2)
    costo_mercancia_total = round(costo_unitario * volumen_unidades, 2)
    gastos_variables_totales = round(gasto_variable_unitario * volumen_unidades, 2)
    utilidad_bruta = round(ingresos_totales - costo_mercancia_total, 2)
    utilidad_neta = round(utilidad_bruta - gastos_variables_totales - gastos_fijos, 2)

    if margen_unitario > 0:
        pe_unidades = math.ceil(gastos_fijos / margen_unitario)
        pe_pesos = round(pe_unidades * precio_efectivo, 2)
    else:
        pe_unidades = None
        pe_pesos = None

    return {
        "precio_lista": round(precio_lista, 2),
        "descuento_pct": round(descuento_pct, 2),
        "precio_efectivo": precio_efectivo,
        "precio_venta_efectivo": precio_efectivo,
        "costo_unitario": round(costo_unitario, 2),
        "costo_total_unitario": costo_total_unitario,
        "gasto_variable_unitario": gasto_variable_unitario,
        "margen_unitario": margen_unitario,
        "margen_porcentual": margen_porcentual,
        "volumen_unidades": round(volumen_unidades, 2),
        "ingresos_totales": ingresos_totales,
        "ingresos_estimados": ingresos_totales,
        "costo_mercancia_total": costo_mercancia_total,
        "utilidad_bruta": utilidad_bruta,
        "utilidad_bruta_estimada": utilidad_bruta,
        "gastos_variables_totales": gastos_variables_totales,
        "gastos_fijos": round(gastos_fijos, 2),
        "utilidad_neta": utilidad_neta,
        "utilidad_neta_estimada": utilidad_neta,
        "punto_equilibrio_unidades": pe_unidades,
        "punto_equilibrio_pesos": pe_pesos,
        "es_rentable": utilidad_neta > 0 and margen_unitario > 0,
    }


@router.post("/simulador")
def simular_escenario_comercial(
    data: SimulacionComercialRequest,
    current_user: models.Usuario = Depends(get_current_user),
):
    """
    Simulador Comercial interactivo (Sección 12):
    Compara la situación actual contra un escenario simulado y entrega las fórmulas
    matemáticas y los supuestos utilizados.
    """
    p_base = (
        data.precio_venta_actual
        if data.precio_venta_actual is not None
        else (data.precio_venta_base if data.precio_venta_base is not None else (data.precio_venta or 0.0))
    )
    c_base = (
        data.costo_compra_actual
        if data.costo_compra_actual is not None
        else (data.costo_compra_base if data.costo_compra_base is not None else (data.precio_compra or 0.0))
    )
    v_base = (
        data.volumen_ventas_actual
        if data.volumen_ventas_actual is not None
        else (data.volumen_base if data.volumen_base is not None else (data.volumen_unidades or 0.0))
    )
    gf_base = (
        data.gastos_fijos
        if data.gastos_fijos is not None
        else data.gastos_fijos_actuales
    )

    p_sim = (
        data.precio_venta_simulado
        if data.precio_venta_simulado is not None
        else (data.precio_venta if data.precio_venta is not None else p_base)
    )
    c_sim = (
        data.costo_compra_simulado
        if data.costo_compra_simulado is not None
        else (data.precio_compra if data.precio_compra is not None else c_base)
    )
    v_sim = (
        data.volumen_ventas_simulado
        if data.volumen_ventas_simulado is not None
        else (data.volumen_unidades if data.volumen_unidades is not None else v_base)
    )
    d_sim = (
        data.descuento_porcentaje
        if data.descuento_porcentaje is not None
        else data.descuento_promocion_pct
    )
    gf_sim = (
        data.gastos_fijos_simulados
        if data.gastos_fijos_simulados is not None
        else gf_base
    )

    actual = _calcular_escenario(
        precio_lista=p_base,
        descuento_pct=0.0,
        costo_unitario=c_base,
        volumen_unidades=v_base,
        gastos_fijos=gf_base,
        gastos_variables_pct=data.gastos_variables_pct_actual,
        gasto_variable_unitario_fijo=data.gasto_variable_unitario,
    )

    simulado = _calcular_escenario(
        precio_lista=p_sim,
        descuento_pct=d_sim,
        costo_unitario=c_sim,
        volumen_unidades=v_sim,
        gastos_fijos=gf_sim,
        gastos_variables_pct=(
            data.gastos_variables_pct_simulado
            if data.gastos_variables_pct_simulado is not None
            else data.gastos_variables_pct_actual
        ),
        gasto_variable_unitario_fijo=data.gasto_variable_unitario,
    )

    dif_utilidad_neta = round(simulado["utilidad_neta"] - actual["utilidad_neta"], 2)
    dif_margen_unitario = round(simulado["margen_unitario"] - actual["margen_unitario"], 2)
    variacion_pct = (
        round((dif_utilidad_neta / abs(actual["utilidad_neta"])) * 100.0, 2)
        if actual["utilidad_neta"] != 0
        else (100.0 if dif_utilidad_neta > 0 else 0.0)
    )

    unidades_para_igualar_actual = None
    if simulado["margen_unitario"] > 0:
        meta_contribucion = actual["utilidad_neta"] + simulado["gastos_fijos"]
        unidades_para_igualar_actual = max(0.0, round(meta_contribucion / simulado["margen_unitario"], 1))

    if simulado["margen_unitario"] <= 0:
        veredicto = "RIESGO_PERDIDA_POR_UNIDAD"
        recomendacion = (
            "Cuidado: En el escenario simulado el precio efectivo no cubre el costo unitario más gastos variables. "
            "Cada unidad vendida generaría pérdida."
        )
    elif dif_utilidad_neta > 0:
        veredicto = "ESCENARIO_FAVORABLE"
        recomendacion = (
            f"El escenario simulado incrementa tu ganancia neta estimada en ${dif_utilidad_neta:,.0f} COP "
            f"({variacion_pct:+.1f}%)."
        )
    elif dif_utilidad_neta < 0:
        veredicto = "ESCENARIO_DESFAVORABLE"
        recomendacion = (
            f"El escenario simulado reduce tu ganancia en ${abs(dif_utilidad_neta):,.0f} COP. "
            f"Necesitarías vender al menos {unidades_para_igualar_actual or 0:g} unidades para mantener tu ganancia actual."
        )
    else:
        veredicto = "NEUTRO"
        recomendacion = "Ambos escenarios producen exactamente la misma ganancia neta estimada."

    bloque_comparacion = {
        "utilidad_neta_base": actual["utilidad_neta"],
        "diferencia_utilidad_neta": dif_utilidad_neta,
        "diferencia_margen_unitario": dif_margen_unitario,
        "variacion_porcentual_utilidad": variacion_pct,
        "unidades_necesarias_igualar_ganancia_actual": unidades_para_igualar_actual,
        "veredicto": veredicto,
        "recomendacion": recomendacion,
    }

    return {
        "nombre_escenario": data.nombre_escenario,
        "situacion_actual": actual,
        "escenario_simulado": simulado,
        "comparacion": bloque_comparacion,
        "comparacion_escenario_base": bloque_comparacion,
        "formulas_documentadas": {
            "precio_efectivo": "Precio Efectivo = Precio de Lista × (1 - % Descuento / 100)",
            "margen_unitario": "Margen por Unidad = Precio Efectivo - Costo de Compra - (Precio Efectivo × % Gasto Variable / 100)",
            "margen_porcentual": "Margen % = (Margen por Unidad / Precio Efectivo) × 100",
            "utilidad_bruta": "Utilidad Bruta = (Precio Efectivo - Costo de Compra) × Volumen de Ventas",
            "utilidad_neta": "Utilidad Neta = Utilidad Bruta - Gastos Variables Totales - Gastos Fijos",
            "punto_equilibrio_unidades": "Punto de Equilibrio (Unidades) = Gastos Fijos / Margen por Unidad (cuando Margen > 0)",
        },
        "supuestos_calculo": [
            "Se asume que todas las unidades proyectadas se venden al precio efectivo del período.",
            "Los costos de compra unitarios se mantienen constantes para el lote proyectado.",
            "Los gastos fijos (arriendo, servicios, nómina base) no dependen del número de unidades vendidas.",
            "Los gastos variables (bolsas, comisiones de pago, transporte por unidad) se expresan como porcentaje del precio efectivo.",
        ],
    }


@router.get("/ventas/excel/{empresa_id}")
def exportar_ventas_excel(
    empresa_id: UUID,
    db: Session = Depends(get_db),
    current_user: models.Usuario = Depends(get_current_user_admin),
):
    if current_user.empresa_id != empresa_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Acceso no autorizado",
        )

    ventas = obtener_ventas_empresa(empresa_id, db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    columnas = [
        "ID Venta", "Fecha", "Estado", "Método Pago", "Vendedor", "Producto",
        "Cantidad", "Precio Unitario", "Subtotal", "Total Venta",
    ]
    ws.append(columnas)

    for celda in ws[1]:
        celda.font = _HEADER_FONT
        celda.fill = _HEADER_FILL
        celda.alignment = Alignment(horizontal="center")

    for venta in ventas:
        for detalle in venta.detalles:
            ws.append([
                str(venta.id),
                venta.fecha_venta.strftime("%Y-%m-%d %H:%M"),
                venta.estado or "completada",
                venta.metodo_pago or "efectivo",
                venta.vendedor_nombre_snapshot or "",
                detalle.producto.nombre if detalle.producto else "",
                float(detalle.cantidad),
                float(detalle.precio_unitario),
                float(detalle.subtotal),
                float(venta.total),
            ])

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=ventas_{empresa_id}.xlsx"
        },
    )


@router.get("/financieros/{empresa_id}")
def obtener_resumen_financiero(
    empresa_id: UUID,
    db: Session = Depends(get_db),
    current_user: models.Usuario = Depends(get_current_user_admin),
):
    if current_user.empresa_id != empresa_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Acceso no autorizado",
        )

    # 1. Inversión en inventario activo (stock * costo)
    inversion_activa = float(
        db.query(
            func.coalesce(func.sum(models.Producto.cantidad_actual * models.Producto.precio_costo), 0)
        )
        .filter(
            models.Producto.empresa_id == empresa_id,
            models.Producto.is_active.is_(True),
        )
        .scalar()
        or 0
    )

    # 2. Ventas totales activas (no anuladas)
    ventas_totales = float(
        db.query(func.coalesce(func.sum(models.Venta.total), 0))
        .filter(
            models.Venta.empresa_id == empresa_id,
            models.Venta.is_active.is_(True),
            models.Venta.estado != "anulada",
        )
        .scalar()
        or 0
    )

    # 2b. Desglose por método de pago (distinguiendo Ventas vs Cobros reales vs Fiado)
    por_metodo_rows = (
        db.query(
            models.Venta.metodo_pago,
            func.coalesce(func.sum(models.Venta.total), 0).label("monto"),
            func.count(models.Venta.id).label("cantidad"),
        )
        .filter(
            models.Venta.empresa_id == empresa_id,
            models.Venta.is_active.is_(True),
            models.Venta.estado != "anulada",
        )
        .group_by(models.Venta.metodo_pago)
        .all()
    )
    desglose_metodos = {
        (r.metodo_pago or "efectivo"): {
            "monto": float(r.monto),
            "cantidad": int(r.cantidad),
        }
        for r in por_metodo_rows
    }
    cuentas_por_cobrar = desglose_metodos.get("fiado", {}).get("monto", 0.0)
    cobros_reales = max(0.0, round(ventas_totales - cuentas_por_cobrar, 2))

    # 3. Costo de productos vendidos (COGS acumulado de ventas válidas)
    cogs_total = float(
        db.query(
            func.coalesce(
                func.sum(
                    models.DetalleVenta.cantidad
                    * func.coalesce(models.DetalleVenta.costo_unitario_snapshot, models.Producto.precio_costo)
                ),
                0,
            )
        )
        .join(models.Venta, models.Venta.id == models.DetalleVenta.venta_id)
        .join(models.Producto, models.DetalleVenta.producto_id == models.Producto.id)
        .filter(
            models.Venta.empresa_id == empresa_id,
            models.Venta.is_active.is_(True),
            models.Venta.estado != "anulada",
        )
        .scalar()
        or 0
    )

    # 4. Ingresos extra y Egresos operativos registrados en caja
    ingresos_extra_caja = float(
        db.query(func.coalesce(func.sum(models.MovimientoCaja.monto), 0))
        .filter(
            models.MovimientoCaja.empresa_id == empresa_id,
            models.MovimientoCaja.is_active.is_(True),
            models.MovimientoCaja.tipo == "ingreso",
        )
        .scalar()
        or 0
    )
    egresos_operativos_caja = float(
        db.query(func.coalesce(func.sum(models.MovimientoCaja.monto), 0))
        .filter(
            models.MovimientoCaja.empresa_id == empresa_id,
            models.MovimientoCaja.is_active.is_(True),
            models.MovimientoCaja.tipo == "egreso",
        )
        .scalar()
        or 0
    )

    utilidad_bruta = round(ventas_totales - cogs_total, 2)
    utilidad_neta_estimada = round(utilidad_bruta + ingresos_extra_caja - egresos_operativos_caja, 2)

    # 5. Costo de reposición para productos con stock bajo
    productos_bajo_stock = (
        db.query(models.Producto)
        .filter(
            models.Producto.empresa_id == empresa_id,
            models.Producto.is_active.is_(True),
            models.Producto.cantidad_actual <= models.Producto.stock_minimo,
        )
        .all()
    )

    costo_reposicion_bajo_stock = sum(
        max(0.0, float(p.stock_minimo or 5.0) * 3.0 - float(p.cantidad_actual)) * float(p.precio_costo)
        for p in productos_bajo_stock
    )

    return {
        # Campos retrocompatibles
        "inversion_activa": round(inversion_activa, 2),
        "ingresos_totales": round(ventas_totales, 2),
        "cogs_total": round(cogs_total, 2),
        "ganancia_neta": utilidad_bruta,
        "costo_reposicion_bajo_stock": round(costo_reposicion_bajo_stock, 2),
        "costo_reposicion_total_sugerido": round(cogs_total + costo_reposicion_bajo_stock, 2),
        # Campos explícitos 360° (Sección 12)
        "ventas_totales": round(ventas_totales, 2),
        "cobros_reales": round(cobros_reales, 2),
        "cuentas_por_cobrar_fiado": round(cuentas_por_cobrar, 2),
        "utilidad_bruta": utilidad_bruta,
        "ingresos_extra_caja": round(ingresos_extra_caja, 2),
        "egresos_operativos_caja": round(egresos_operativos_caja, 2),
        "utilidad_neta_estimada": utilidad_neta_estimada,
        "desglose_metodos_pago": desglose_metodos,
    }
